#!/usr/bin/env python3
"""
Scalify Agency - AI Outbound Calling System
===========================================
Places outbound calls from your SignalWire number and bridges each answered
call to the AI agent (Jessica) over SIP using inline TwiML/LaML (the audio fix).

Commands
  python main.py call <number> [--name NAME] [--dry-run]
  python main.py bulk <leads.csv|leads.xlsx> [--max-concurrent N] [--delay S] [--limit N] [--dry-run]
  python main.py status <CALL_SID>
  python main.py sync        # refresh status / duration / recording URL for logged calls

Safety
  TEST_MODE=true in .env (the default) redirects EVERY call to TEST_NUMBER,
  forces bulk runs to one call at a time, and caps them at 3 calls unless
  --limit is given.

Output
  logs/scalify_calls.log   human-readable app log (rotating)
  logs/calls.csv           one row per call: SID, status, duration, recording URL
"""
from __future__ import annotations

import argparse
import csv
import logging
import os
import re
import sys
import time
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from logging.handlers import RotatingFileHandler
from pathlib import Path
from xml.sax.saxutils import escape

try:
    from dotenv import load_dotenv
except ImportError:  # dotenv is optional if env vars are set another way
    load_dotenv = None

# --------------------------------------------------------------------------- #
# Constants
# --------------------------------------------------------------------------- #
BASE_DIR = Path(__file__).resolve().parent
LOG_DIR = BASE_DIR / "logs"
CALL_LOG = LOG_DIR / "calls.csv"

ACTIVE_STATUSES = {"queued", "initiated", "ringing", "in-progress"}
TERMINAL_STATUSES = {"completed", "busy", "failed", "no-answer", "canceled"}
RETRYABLE_HTTP = {429, 500, 502, 503, 504}
MAX_RETRIES = 3
POLL_SECONDS = 5
MAX_TRACK_SECONDS = 30 * 60  # stop tracking a call after 30 min (sync picks it up later)
TEST_MODE_BULK_CAP = 3

LOG_FIELDS = [
    "created_at", "lead_name", "lead_phone", "dialed_to", "call_sid",
    "status", "duration_sec", "recording_url", "test_mode", "error",
]

log = logging.getLogger("scalify")


class ConfigError(Exception):
    """Raised when .env / environment configuration is missing or invalid."""


# --------------------------------------------------------------------------- #
# Logging
# --------------------------------------------------------------------------- #
def setup_logging(verbose: bool = False) -> None:
    LOG_DIR.mkdir(exist_ok=True)
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%Y-%m-%d %H:%M:%S")

    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(fmt)
    file_handler = RotatingFileHandler(
        LOG_DIR / "scalify_calls.log", maxBytes=5_000_000, backupCount=5, encoding="utf-8"
    )
    file_handler.setFormatter(fmt)

    log.setLevel(logging.DEBUG if verbose else logging.INFO)
    log.handlers[:] = [console, file_handler]
    log.propagate = False
    logging.getLogger("twilio").setLevel(logging.WARNING)  # SDK is chatty


# --------------------------------------------------------------------------- #
# Phone numbers
# --------------------------------------------------------------------------- #
def normalize_us_number(raw: str) -> str:
    """Return a US/NANP number in E.164 (+1XXXXXXXXXX) or raise ValueError."""
    digits = re.sub(r"\D", "", str(raw or ""))
    if len(digits) == 10:
        digits = "1" + digits
    if len(digits) != 11 or not digits.startswith("1"):
        raise ValueError(f"not a valid US number: {raw!r}")
    # NANP rule: area code and exchange can't start with 0 or 1
    if digits[1] in "01" or digits[4] in "01":
        raise ValueError(f"invalid US area code/exchange: {raw!r}")
    return "+" + digits


# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
def _env_bool(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None or not value.strip():
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


@dataclass(frozen=True)
class Config:
    project_id: str
    api_token: str
    space_host: str
    from_number: str
    agent_sip: str
    test_mode: bool
    test_number: str
    status_callback_url: str
    record_calls: bool
    ring_timeout: int

    @classmethod
    def load(cls, require_credentials: bool = True) -> "Config":
        if load_dotenv:
            load_dotenv(BASE_DIR / ".env")

        # Accept both SIGNALWIRE_* names and the short names from your old .env
        project_id = (os.getenv("SIGNALWIRE_PROJECT_ID") or os.getenv("PROJECT_ID") or "").strip()
        api_token = (os.getenv("SIGNALWIRE_API_TOKEN") or os.getenv("API_TOKEN") or "").strip()
        if require_credentials and not (project_id and api_token):
            raise ConfigError("SIGNALWIRE_PROJECT_ID and SIGNALWIRE_API_TOKEN must be set in .env")

        space = (os.getenv("SIGNALWIRE_SPACE_URL") or os.getenv("SPACE_URL") or "").strip()
        space_host = re.sub(r"^https?://", "", space).rstrip("/")
        if not space_host.endswith(".signalwire.com"):
            raise ConfigError("SIGNALWIRE_SPACE_URL must look like yourspace.signalwire.com")

        agent_sip = os.getenv("AGENT_SIP", "").strip()
        if not agent_sip.startswith("sip:"):
            raise ConfigError("AGENT_SIP must be a SIP URI starting with 'sip:'")

        try:
            from_number = normalize_us_number(os.getenv("FROM_NUMBER", ""))
            test_number = normalize_us_number(os.getenv("TEST_NUMBER", ""))
        except ValueError as exc:
            raise ConfigError(f"FROM_NUMBER / TEST_NUMBER: {exc}") from exc

        try:
            ring_timeout = int(os.getenv("RING_TIMEOUT", "30"))
        except ValueError as exc:
            raise ConfigError("RING_TIMEOUT must be an integer (seconds)") from exc

        return cls(
            project_id=project_id,
            api_token=api_token,
            space_host=space_host,
            from_number=from_number,
            agent_sip=agent_sip,
            test_mode=_env_bool("TEST_MODE", True),  # safe default: ON
            test_number=test_number,
            status_callback_url=os.getenv("STATUS_CALLBACK_URL", "").strip(),
            record_calls=_env_bool("RECORD_CALLS", True),
            ring_timeout=max(5, min(ring_timeout, 120)),
        )


# --------------------------------------------------------------------------- #
# SignalWire client + TwiML
# --------------------------------------------------------------------------- #
def get_client(cfg: Config):
    try:
        from signalwire.rest import Client
    except ImportError as exc:
        raise ConfigError("SignalWire SDK missing - run: pip install -r requirements.txt") from exc
    return Client(cfg.project_id, cfg.api_token, signalwire_space_url=cfg.space_host)


def build_twiml(
    cfg: Config,
    lead_name: str = "",
    business_name: str = "",
    lead_source: str = "",
    lead_phone: str = "",
) -> str:
    """Inline TwiML that bridges the answered call to Jessica over SIP.

    Using inline `twiml=` instead of a `url=` is what fixed the audio issue.
    Optional SIP headers carry lead context to the local agent (Option 2).
    """
    attrs = ' record="record-from-answer-dual"' if cfg.record_calls else ""
    headers = ""
    for name, value in (
        ("X-Lead-Name", lead_name),
        ("X-Business-Name", business_name),
        ("X-Lead-Source", lead_source),
        ("X-Lead-Phone", lead_phone),
    ):
        if value:
            headers += f'<Header name="{escape(name)}" value="{escape(value)}"/>'
    return (
        f"<Response><Dial{attrs}><Sip>{headers}{escape(cfg.agent_sip)}</Sip></Dial></Response>"
    )


# --------------------------------------------------------------------------- #
# Call records / CSV log
# --------------------------------------------------------------------------- #
def now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


@dataclass
class CallRecord:
    created_at: str = ""
    lead_name: str = ""
    lead_phone: str = ""
    dialed_to: str = ""
    call_sid: str = ""
    status: str = ""
    duration_sec: str = ""
    recording_url: str = ""
    test_mode: str = ""
    error: str = ""


def append_log(rec: CallRecord) -> None:
    LOG_DIR.mkdir(exist_ok=True)
    is_new = not CALL_LOG.exists()
    with CALL_LOG.open("a", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=LOG_FIELDS)
        if is_new:
            writer.writeheader()
        writer.writerow(asdict(rec))


def read_log() -> list[dict]:
    if not CALL_LOG.exists():
        return []
    with CALL_LOG.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def write_log(rows: list[dict]) -> None:
    tmp = CALL_LOG.with_suffix(".tmp")
    with tmp.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=LOG_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    tmp.replace(CALL_LOG)


def update_log(call_sid: str, **fields: str) -> None:
    rows = read_log()
    for row in rows:
        if row.get("call_sid") == call_sid:
            row.update({k: v for k, v in fields.items() if v not in (None, "")})
    write_log(rows)


# --------------------------------------------------------------------------- #
# Core: place a call
# --------------------------------------------------------------------------- #
def place_call(
    client,
    cfg: Config,
    lead_phone: str,
    lead_name: str = "",
    dry_run: bool = False,
    business_name: str = "",
    lead_source: str = "",
) -> CallRecord:
    rec = CallRecord(created_at=now_iso(), lead_name=lead_name, lead_phone=str(lead_phone),
                     test_mode=str(cfg.test_mode).lower())
    try:
        target = normalize_us_number(lead_phone)
    except ValueError as exc:
        rec.status, rec.error = "invalid_number", str(exc)
        log.warning("Skipping %s: %s", lead_name or "lead", exc)
        return rec

    rec.lead_phone = target
    rec.dialed_to = cfg.test_number if cfg.test_mode else target
    if rec.dialed_to == cfg.from_number:
        rec.status, rec.error = "invalid_number", "destination equals FROM_NUMBER"
        log.warning("Skipping %s: destination equals FROM_NUMBER", target)
        return rec

    params = {
        "to": rec.dialed_to,
        "from_": cfg.from_number,
        "twiml": build_twiml(
            cfg,
            lead_name=lead_name,
            business_name=business_name,
            lead_source=lead_source,
            lead_phone=target,
        ),
        "timeout": cfg.ring_timeout,
    }
    if cfg.status_callback_url:
        params.update(
            status_callback=cfg.status_callback_url,
            status_callback_event=["initiated", "ringing", "answered", "completed"],
            status_callback_method="POST",
        )

    label = f"{lead_name or 'lead'} {target}"
    if cfg.test_mode and target != rec.dialed_to:
        label += f" (TEST MODE -> {rec.dialed_to})"

    if dry_run:
        rec.status = "dry_run"
        log.info("[DRY RUN] Would call %s | TwiML: %s", label, params["twiml"])
        return rec

    import requests
    from twilio.base.exceptions import TwilioRestException

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            call = client.calls.create(**params)
            rec.call_sid = call.sid
            rec.status = call.status or "queued"
            rec.error = ""
            log.info("Call placed: %s | SID %s | %s", label, call.sid, rec.status)
            return rec
        except TwilioRestException as exc:
            retry = exc.status in RETRYABLE_HTTP and attempt < MAX_RETRIES
            rec.error = f"HTTP {exc.status} / code {exc.code}: {exc.msg}"
            log.log(logging.WARNING if retry else logging.ERROR,
                    "SignalWire API error for %s (attempt %d/%d): %s",
                    label, attempt, MAX_RETRIES, rec.error)
            if not retry:
                break
        except requests.exceptions.RequestException as exc:
            retry = attempt < MAX_RETRIES
            rec.error = f"network error: {exc}"
            log.warning("Network error for %s (attempt %d/%d): %s", label, attempt, MAX_RETRIES, exc)
            if not retry:
                break
        time.sleep(2 ** attempt)  # 2s, 4s backoff

    rec.status = "api_error"
    return rec


# --------------------------------------------------------------------------- #
# Call details (status, duration, recording)
# --------------------------------------------------------------------------- #
def _recording_url(cfg: Config, recording_sid: str) -> str:
    return (f"https://{cfg.space_host}/api/laml/2010-04-01/Accounts/"
            f"{cfg.project_id}/Recordings/{recording_sid}.mp3")


def fetch_call_details(client, cfg: Config, call_sid: str) -> dict:
    call = client.calls(call_sid).fetch()
    details = {
        "status": call.status or "",
        "duration_sec": str(call.duration or ""),
        "recording_url": "",
    }
    if cfg.record_calls and call.status == "completed":
        try:
            recs = client.recordings.list(call_sid=call_sid, limit=1)
            if not recs:  # <Dial record> can attach to the bridged (child) leg
                for child in client.calls.list(parent_call_sid=call_sid, limit=5):
                    recs = client.recordings.list(call_sid=child.sid, limit=1)
                    if recs:
                        break
            if recs:
                details["recording_url"] = _recording_url(cfg, recs[0].sid)
        except Exception as exc:  # recording lookup must never break call tracking
            log.debug("Recording lookup failed for %s: %s", call_sid, exc)
    return details


# --------------------------------------------------------------------------- #
# Bulk calling
# --------------------------------------------------------------------------- #
PHONE_KEYS = ("phone", "phone_number", "phone number", "number", "mobile", "cell")


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 7869848504 as 7869848504.0
    return str(value).strip()


def load_leads(path: str | Path) -> list[dict]:
    """Read leads from .csv or .xlsx. Needs a phone column; name is optional."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"lead file not found: {path}")

    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            raw_rows = list(csv.DictReader(fh))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConfigError("openpyxl needed for Excel files: pip install openpyxl") from exc
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [_cell_to_str(h) for h in next(rows_iter, [])]
        raw_rows = [dict(zip(headers, map(_cell_to_str, row))) for row in rows_iter]
    else:
        raise ValueError("lead file must be .csv or .xlsx")

    leads, seen = [], set()
    for row in raw_rows:
        row = {str(k).strip().lower(): _cell_to_str(v) for k, v in row.items() if k}
        phone_key = next((k for k in PHONE_KEYS if k in row), None)
        if phone_key is None:
            raise ValueError(f"no phone column found; expected one of {PHONE_KEYS}")
        phone = row.get(phone_key, "")
        if not phone:
            continue
        name = row.get("name") or " ".join(
            p for p in (row.get("first_name", ""), row.get("last_name", "")) if p
        )
        dedupe_key = re.sub(r"\D", "", phone)[-10:]
        if dedupe_key in seen:
            log.info("Duplicate number skipped: %s", phone)
            continue
        seen.add(dedupe_key)
        leads.append({"phone": phone, "name": name, **row})
    return leads


def _refresh_active(client, cfg: Config, active: dict, summary: Counter) -> None:
    for sid in list(active):
        started = active[sid]
        try:
            details = fetch_call_details(client, cfg, sid)
        except Exception as exc:
            log.warning("Could not fetch status for %s: %s", sid, exc)
            continue
        status = details["status"]
        if status in TERMINAL_STATUSES:
            update_log(sid, **details)
            summary[status] += 1
            log.info("Call %s finished: %s (%ss)", sid, status, details["duration_sec"] or 0)
            active.pop(sid)
        elif time.monotonic() - started > MAX_TRACK_SECONDS:
            log.warning("Call %s still '%s' after 30 min - no longer tracking (run sync later)", sid, status)
            update_log(sid, status=status)
            active.pop(sid)


def run_bulk(client, cfg: Config, leads: list[dict], max_concurrent: int = 1,
             delay: float = 10, limit: int | None = None, dry_run: bool = False) -> Counter:
    if cfg.test_mode:
        if max_concurrent > 1:
            log.warning("TEST MODE: all calls go to one number - forcing max-concurrent=1")
            max_concurrent = 1
        if limit is None:
            limit = TEST_MODE_BULK_CAP
            log.warning("TEST MODE: capping bulk run at %d calls (use --limit to change)", limit)
    if limit:
        leads = leads[:limit]

    log.info("Bulk run: %d lead(s), max-concurrent=%d, delay=%ss%s",
             len(leads), max_concurrent, delay, " [DRY RUN]" if dry_run else "")

    active: dict[str, float] = {}  # call_sid -> monotonic start time
    summary: Counter = Counter()
    try:
        for i, lead in enumerate(leads, 1):
            while not dry_run and len(active) >= max_concurrent:
                time.sleep(POLL_SECONDS)
                _refresh_active(client, cfg, active, summary)

            log.info("[%d/%d] %s", i, len(leads), lead.get("name") or lead["phone"])
            rec = place_call(
                client,
                cfg,
                lead["phone"],
                lead.get("name", ""),
                dry_run=dry_run,
                business_name=lead.get("company") or lead.get("business_name") or lead.get("business") or "",
                lead_source=lead.get("lead_source") or lead.get("source") or "",
            )
            if not dry_run:
                append_log(rec)
            if rec.call_sid:
                active[rec.call_sid] = time.monotonic()
            else:
                summary[rec.status] += 1

            if not dry_run and delay > 0 and i < len(leads):
                time.sleep(delay)

        while active:
            time.sleep(POLL_SECONDS)
            _refresh_active(client, cfg, active, summary)
    except KeyboardInterrupt:
        log.warning("Interrupted - no new calls will be placed. %d call(s) may still be live; "
                    "run `python main.py sync` later to update the log.", len(active))

    log.info("Bulk run finished: %s", dict(summary) or "no calls")
    return summary


# --------------------------------------------------------------------------- #
# Sync / status commands
# --------------------------------------------------------------------------- #
def sync_log(client, cfg: Config) -> int:
    rows = read_log()
    updated = 0
    for row in rows:
        sid = row.get("call_sid")
        if not sid:
            continue
        needs_status = row.get("status") not in TERMINAL_STATUSES
        needs_recording = (cfg.record_calls and row.get("status") == "completed"
                           and not row.get("recording_url"))
        if not (needs_status or needs_recording or not row.get("duration_sec")):
            continue
        try:
            details = fetch_call_details(client, cfg, sid)
        except Exception as exc:
            log.warning("Could not fetch %s: %s", sid, exc)
            continue
        row.update({k: v for k, v in details.items() if v})
        updated += 1
    write_log(rows)
    log.info("Synced %d call(s) in %s", updated, CALL_LOG)
    return updated


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Scalify Agency AI outbound calling (agent: Jessica)")
    parser.add_argument("-v", "--verbose", action="store_true", help="debug logging")
    sub = parser.add_subparsers(dest="command", required=True)

    p_call = sub.add_parser("call", help="call one number")
    p_call.add_argument("number")
    p_call.add_argument("--name", default="")
    p_call.add_argument("--business", default="", help="business / company name for Jessica")
    p_call.add_argument("--source", default="", help="lead source for Jessica")
    p_call.add_argument("--dry-run", action="store_true", help="show what would happen, place no call")

    p_bulk = sub.add_parser("bulk", help="call leads from a CSV/XLSX file")
    p_bulk.add_argument("file")
    p_bulk.add_argument("--max-concurrent", type=int, default=1)
    p_bulk.add_argument("--delay", type=float, default=10, help="seconds between new calls")
    p_bulk.add_argument("--limit", type=int, default=None)
    p_bulk.add_argument("--dry-run", action="store_true")

    p_status = sub.add_parser("status", help="show details for one call SID")
    p_status.add_argument("call_sid")

    sub.add_parser("sync", help="refresh status/duration/recording for logged calls")
    sub.add_parser("serve-agent", help="run local Jessica agent (SWML/SWAIG) on AGENT_PORT")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    setup_logging(args.verbose)
    dry_run = getattr(args, "dry_run", False)

    if args.command == "serve-agent":
        try:
            from agent.jessica_agent import build_agent
        except Exception:
            log.exception("Failed to import Jessica agent")
            return 1
        try:
            agent = build_agent()
            log.info(
                "Serving Jessica on port %s — set SWML_PROXY_URL_BASE to your ngrok URL",
                os.getenv("AGENT_PORT", "3000"),
            )
            agent.run(
                host=os.getenv("AGENT_HOST", "0.0.0.0"),
                port=int(os.getenv("AGENT_PORT", "3000")),
            )
            return 0
        except Exception as exc:
            log.error("%s", exc)
            return 1

    try:
        cfg = Config.load(require_credentials=not dry_run)
        client = None if dry_run else get_client(cfg)
    except ConfigError as exc:
        log.error("Config error: %s", exc)
        return 2

    mode = f"TEST (all calls -> {cfg.test_number})" if cfg.test_mode else "LIVE"
    log.info("Mode: %s | From: %s | Recording: %s", mode, cfg.from_number,
             "on" if cfg.record_calls else "off")

    try:
        if args.command == "call":
            rec = place_call(
                client,
                cfg,
                args.number,
                args.name,
                dry_run=dry_run,
                business_name=args.business,
                lead_source=args.source,
            )
            if not dry_run:
                append_log(rec)
            return 0 if rec.call_sid or dry_run else 1

        if args.command == "bulk":
            leads = load_leads(args.file)
            if not leads:
                log.error("No leads with phone numbers in %s", args.file)
                return 1
            run_bulk(client, cfg, leads, max(1, args.max_concurrent), max(0.0, args.delay),
                     args.limit, dry_run=dry_run)
            return 0

        if args.command == "status":
            details = fetch_call_details(client, cfg, args.call_sid)
            for key, value in details.items():
                print(f"{key:>14}: {value or '-'}")
            if details["status"] in TERMINAL_STATUSES:
                update_log(args.call_sid, **details)
            return 0

        if args.command == "sync":
            sync_log(client, cfg)
            return 0
    except (FileNotFoundError, ValueError, ConfigError) as exc:
        log.error("%s", exc)
        return 1
    except Exception:
        log.exception("Unexpected error")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
